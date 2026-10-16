"""Inspect the Simplified list asset source without creating an import register.

Requires openpyxl. Never evaluates formulas, assigns IDs/levels or contacts a DB.
The JSON report contains private source values and belongs in ignored output/.
"""
import argparse
from collections import Counter, defaultdict
import datetime as dt
import hashlib
import json
from pathlib import Path
import zipfile

import openpyxl

CORE = (
    "Commissionable Asset", "Discipline", "Location / Area", "Product/System",
    "Design phase", "FR Zone", "(QCx L1-4 Asset Type)",
)
REQUESTED_ABSENT = (
    "FieldTrack internal ID", "Client asset ID", "Schedule ID / Activity ID mapping",
    "Floor", "IST group membership", "Required commissioning levels",
    "Shared group test-sheet rules",
)


def blank(value):
    return value is None or isinstance(value, str) and not value.strip()


def json_value(value):
    if isinstance(value, (dt.date, dt.datetime, dt.time)):
        return {"type": type(value).__name__, "iso": value.isoformat()}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def classify_rows(headers, rows):
    """rows contains (source row number, cached values, formula metadata)."""
    if len(set(headers)) != len(headers) or any(not name for name in headers):
        raise ValueError("Source headers must be nonblank and unique.")
    if any(name not in headers for name in CORE):
        raise ValueError("Expected all seven Simplified list core headers.")
    candidates, excluded, missing = [], [], Counter()
    name_groups, exact_groups = defaultdict(list), defaultdict(list)
    blank_rows, formula_count, missing_cache_count = 0, 0, 0
    distributions = {name: Counter() for name in CORE[1:]}
    for number, values, formulas in rows:
        values = list(values)
        if len(values) != len(headers):
            raise ValueError(f"Row {number} does not match the source columns.")
        data = dict(zip(headers, map(json_value, values)))
        formula_count += len(formulas)
        missing_cache_count += sum(item["cachedValue"] is None for item in formulas)
        if all(blank(value) for value in values) and not formulas:
            blank_rows += 1
            continue
        entry = {"sourceRow": number, "sourceValues": data, "formulas": formulas}
        # A free-text name/notes-only row is retained for review, never promoted to an asset.
        if not any(not blank(data[name]) for name in CORE[1:]):
            entry["reason"] = "No populated core classification fields; review note/unclassified row."
            excluded.append(entry)
            continue
        absent = [name for name in CORE if blank(data[name])]
        entry["missingCoreFields"] = absent
        entry["missingSourceFields"] = [name for name, value in data.items() if blank(value)]
        entry["requiresReview"] = bool(absent or formulas)
        for name in entry["missingSourceFields"]:
            missing[name] += 1
        for name, distribution in distributions.items():
            value = data[name]
            if not blank(value):
                distribution[str(value).strip()] += 1
        name = data[CORE[0]]
        if not blank(name):
            name_groups[str(name).strip()].append(number)
        exact_groups[json.dumps(data, sort_keys=True, ensure_ascii=False)].append(number)
        candidates.append(entry)
    duplicate_names = [
        {"nameAfterTrim": name, "sourceRows": numbers}
        for name, numbers in sorted(name_groups.items()) if len(numbers) > 1
    ]
    exact_duplicates = [numbers for numbers in exact_groups.values() if len(numbers) > 1]
    duplicate_rows = {row for group in duplicate_names for row in group["sourceRows"]}
    for entry in candidates:
        entry["duplicateName"] = entry["sourceRow"] in duplicate_rows
        entry["requiresReview"] |= entry["duplicateName"]
    return {
        "summary": {
            "candidateRows": len(candidates),
            "distinctTrimmedNames": len(name_groups),
            "duplicateNameGroups": len(duplicate_names),
            "excessRowsByName": sum(len(group["sourceRows"]) - 1 for group in duplicate_names),
            "exactDuplicateGroups": len(exact_duplicates),
            "excludedNonAssetRows": len(excluded),
            "blankRows": blank_rows,
            "rowsWithMissingCoreFields": sum(bool(row["missingCoreFields"]) for row in candidates),
            "formulaCells": formula_count,
            "formulaCellsWithoutCachedValue": missing_cache_count,
            "missingSourceFields": dict(sorted(missing.items())),
        },
        "distributions": {name: dict(sorted(counts.items())) for name, counts in distributions.items()},
        "duplicateNames": duplicate_names,
        "exactDuplicateSourceRows": exact_duplicates,
        "excludedRows": excluded,
        "candidateRows": candidates,
    }


def inspect_workbook(path, sheet_name="Simplified list"):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Expected an .xlsx source workbook.")
    original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    with zipfile.ZipFile(path) as package:
        if sum(part.file_size for part in package.infolist()) > 100_000_000:
            raise ValueError("Workbook exceeds the 100 MB uncompressed inspection limit.")
        parts = package.namelist()
        if any(name.lower().endswith("vbaproject.bin") for name in parts):
            raise ValueError("Macro content is outside this XLSX preflight.")
        external_links = any(name.startswith("xl/externalLinks/") for name in parts)
    source = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if sheet_name not in source.sheetnames:
            raise ValueError("Expected the explicitly selected source sheet.")
        sheet, cache = source[sheet_name], cached[sheet_name]
        # Google Sheets exports may omit worksheet dimension metadata.
        # Compute it from the read-only XML rather than guessing the data extent.
        if sheet.max_row is None or sheet.max_column is None:
            sheet.calculate_dimension(force=True)
        if cache.max_row is None or cache.max_column is None:
            cache.calculate_dimension(force=True)
        if sheet.max_row > 100_000 or sheet.max_column > 256:
            raise ValueError("Source exceeds the 100,000 row / 256 column inspection limit.")
        headers = [str(cell.value).strip() if cell.value is not None else ""
                   for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

        def rows():
            originals = sheet.iter_rows(min_row=2)
            caches = cache.iter_rows(min_row=2)
            for number, (raw, values) in enumerate(zip(originals, caches, strict=True), 2):
                formulas = []
                for cell, value in zip(raw, values, strict=True):
                    if cell.data_type == "f":
                        formulas.append({"cell": cell.coordinate, "formula": json_value(cell.value),
                                         "cachedValue": json_value(value.value),
                                         "cacheVerified": False})
                yield number, [cell.value for cell in values], formulas

        result = classify_rows(headers, rows())
        result = {
            "format": "fieldtrack-asset-preflight-v1",
            "reviewOnly": True,
            "importAllowed": False,
            "source": {"fileName": path.name, "sha256": original_hash,
                       "bytes": path.stat().st_size, "sheet": sheet_name,
                       "headerRow": 1, "physicalLastRow": sheet.max_row,
                       "headers": headers, "externalLinksPresent": external_links},
            "policy": {
                "identityGenerated": False, "levelsInferred": False,
                "deduplicationApplied": False, "formulasEvaluated": False,
                "databaseAccess": False,
                "candidateRule": "At least one populated core classification field (core fields except name). Missing core fields are retained for review.",
                "duplicateRule": "Case-sensitive asset name after trim; collisions remain separate source rows.",
                "formulaPolicy": "Formula text and saved cache retained separately. Cached results are unverified and never recalculated.",
                "unresolvedInputs": list(REQUESTED_ABSENT),
            },
            **result,
        }
    finally:
        source.close()
        if cached is not None:
            cached.close()
    if hashlib.sha256(path.read_bytes()).hexdigest() != original_hash:
        raise ValueError("Source changed during inspection. Retry against a stable copy.")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path, help="Private .review.json report; never an import payload")
    parser.add_argument("--sheet", default="Simplified list")
    args = parser.parse_args()
    if not args.output.name.endswith(".review.json"):
        parser.error("Output must end in .review.json.")
    if args.output.resolve() == args.workbook.resolve() or args.output.exists():
        parser.error("Output must be a new file and cannot overwrite the source or an existing report.")
    try:
        report = inspect_workbook(args.workbook, args.sheet)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open("x", encoding="utf-8") as output:
            json.dump(report, output, ensure_ascii=False, indent=2, allow_nan=False)
            output.write("\n")
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        parser.exit(1, f"Preflight failed: {error}\n")
    print(json.dumps({"reviewOnly": True, "importAllowed": False, **report["summary"]}))


if __name__ == "__main__":
    main()
