"""Read-only conversion of Tony's P6 / AMS16 workbooks to an import register.

No network/database writes. Ambiguous data is reported alongside the output.
Asset levels are explicit: --levels L2A L2B L3. No completion inferred from
mixed Actual / Forecast columns or coloured cells.
"""
import argparse
import datetime as dt
import json
import re
from pathlib import Path
import openpyxl


def date_value(value):
    if value is None or str(value).strip() in ('', 'N/A', '-', 'TBC'):
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    text = str(value).strip().rstrip('*').strip()
    if text.endswith(' A'):
        text = text[:-2].strip()
    for pattern in ('%d-%b-%y', '%d-%b-%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return dt.datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f'Unrecognized date: {value!r}')


def convert(path, levels):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    assets, schedule, warnings = {}, {}, []
    def activity(activity_id, name, start, finish, source):
        row = {'activityId': str(activity_id).strip(), 'name': name, 'source': source}
        for key, raw in [('start', start), ('finish', finish)]:
            try:
                value = date_value(raw)
                if value:
                    row[key] = value
            except ValueError as error:
                warnings.append(f'{source}: {error}')
        if row.get('start') and row.get('finish') and row['start'] > row['finish']:
            warnings.append(f'{source}: finish before start; omitted dates for review')
            row.pop('start', None)
            row.pop('finish', None)
        old = schedule.get(row['activityId'])
        if old:
            if any(old.get(k) != row.get(k) for k in ['start', 'finish']):
                warnings.append(f"{source}: conflicting duplicate {row['activityId']}; first occurrence retained")
        else:
            schedule[row['activityId']] = row
    if 'Tagging Schedule' in workbook.sheetnames:
        sheet = workbook['Tagging Schedule']
        for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            if number <= 3 or not cells[1] or not cells[2]:
                continue  # WBS/summary rows do not have an activity name.
            activity(cells[1], str(cells[2]).strip(), cells[4], cells[5], f'{path.name} / {sheet.title} / row {number}')
    elif 'Cx Plan' in workbook.sheetnames:
        sheet = workbook['Cx Plan']
        rows = sheet.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else '' for v in next(rows)]
        id_columns = [i for i, h in enumerate(headers) if h == 'Activity ID']
        for number, cells in enumerate(rows, 2):
            asset_id = str(cells[0] or '').strip()
            if not asset_id.startswith('AMS16_'):
                continue
            name = str(cells[6] or asset_id).strip()
            equipment_type = re.sub(r'_S\d+.*$', '', name).strip()
            if asset_id in assets:
                warnings.append(f'{sheet.title} row {number}: duplicate asset {asset_id}; first occurrence retained')
                continue
            assets[asset_id] = {'id': asset_id, 'name': name, 'equipmentType': equipment_type, 'location': str(cells[8] or '').strip(), 'levels': levels}
            for col in id_columns:
                activity_id = str(cells[col] or '').strip()
                if not activity_id or activity_id == 'N/A':
                    continue
                phase = headers[col+1] if col+1 < len(headers) else 'Commissioning'
                # P6 planned date has explicit provenance. Actual / Forecast columns are intentionally not used as actuals.
                activity(activity_id, f'{asset_id} — {phase}', None, cells[col+1], f'{path.name} / {sheet.title} / row {number} / {openpyxl.utils.get_column_letter(col+1)}')
        warnings.append('Asset levels are selected by the operator, not inferred from tag colour. Verify exceptions per asset before applying.')
        warnings.append('Task-to-asset/level/Activity ID links require explicit mapping. Actual/Forecast fields and checklist statuses were not converted to approvals.')
    else:
        raise ValueError('Expected Tagging Schedule (P6) or Cx Plan (AMS16).')
    return {'assets': list(assets.values()), 'schedule': list(schedule.values())}, warnings


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--levels', nargs='*', choices=['L1','L2','L2A','L2B','L3','L4','L5','L6'], default=[])
    args = parser.parse_args()
    result, warnings = convert(args.workbook, list(dict.fromkeys(args.levels)))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False))
    args.output.with_suffix('.review.json').write_text(json.dumps({'source': args.workbook.name, 'warnings': warnings}, indent=2, ensure_ascii=False))
    print(f"{len(result['assets'])} assets, {len(result['schedule'])} schedule activities, {len(warnings)} review notes. Source unchanged.")


if __name__ == '__main__':
    main()
